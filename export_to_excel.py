#!/usr/bin/env python3
"""
Export ai_events.db to a multi-sheet ai_events.xlsx workbook.

One sheet per table found in the database: AI_Event, sources_raw, plus every
enrichment table added alongside the 15 new sources (Vulnerability_Report,
Risk_Taxonomy, Risk_Register, Atlas_Taxonomy, Legal_Case, Security_Advisory,
InspectAgents_Report). A table with zero rows still gets a sheet with just
headers (empty is expected/fine for fallback-only sources with no file
supplied).

Run with:  python3 export_to_excel.py [--db ai_events.db] [--out ai_events.xlsx]
"""

import argparse
import sqlite3

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Preferred sheet order; any other table present in the DB is appended after these.
PREFERRED_ORDER = [
    "AI_Event",
    "sources_raw",
    "Vulnerability_Report",
    "Risk_Taxonomy",
    "Risk_Register",
    "Atlas_Taxonomy",
    "Legal_Case",
    "Security_Advisory",
    "InspectAgents_Report",
]

MAX_CELL_LEN = 32000  # Excel cell limit is ~32767; raw_json can be huge


def list_tables(conn):
    cur = conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
    )
    return [r[0] for r in cur.fetchall()]


def export(db_path, out_path):
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    tables = list_tables(conn)
    ordered = [t for t in PREFERRED_ORDER if t in tables]
    ordered += [t for t in tables if t not in ordered]

    wb = Workbook()
    wb.remove(wb.active)

    for table in ordered:
        ws = wb.create_sheet(title=table[:31])  # Excel sheet-name length limit
        cur = conn.execute(f"SELECT * FROM {table}")
        columns = [d[0] for d in cur.description]
        ws.append(columns)
        row_count = 0
        for row in cur:
            values = []
            for v in row:
                if isinstance(v, str) and len(v) > MAX_CELL_LEN:
                    v = v[:MAX_CELL_LEN] + " …[truncated]"
                values.append(v)
            ws.append(values)
            row_count += 1


        # light header styling + reasonable column widths
        ws.views.sheetView[0].showGridLines = True
        bold_font = Font(bold=True)
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=1, column=col_idx).font = ws.cell(row=1, column=col_idx).font.copy(bold=True)
            width = min(max(len(col_name) + 2, 12), 40)
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"
        print(f"  {table:24s}: {row_count} rows")

    wb.save(out_path)
    conn.close()
    print(f"Wrote {out_path} with {len(ordered)} sheet(s).")


def parse_args():
    p = argparse.ArgumentParser(description="Export ai_events.db to a multi-sheet xlsx workbook")
    p.add_argument("--db", default="ai_events.db", help="Path to the SQLite database")
    p.add_argument("--out", default="ai_events.xlsx", help="Path to write the .xlsx workbook")
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    export(args.db, args.out)
